"""Word (.docx) ve PDF dosyalarindan duz metin cikarir."""
import re
from io import BytesIO
from pathlib import Path


def word_oku(veri: bytes) -> str:
    from docx import Document

    belge = Document(BytesIO(veri))
    satirlar = [p.text.strip() for p in belge.paragraphs if p.text.strip()]

    # Tablolardaki metni de al (satir satir, hucreler " | " ile)
    for tablo in belge.tables:
        for satir in tablo.rows:
            hucreler = [h.text.strip() for h in satir.cells if h.text.strip()]
            if hucreler:
                satirlar.append(" | ".join(hucreler))

    return "\n".join(satirlar)


def pdf_oku(veri: bytes) -> str:
    import pdfplumber

    sayfalar = []
    with pdfplumber.open(BytesIO(veri)) as pdf:
        for sayfa in pdf.pages:
            metin = sayfa.extract_text() or ""
            if metin.strip():
                sayfalar.append(metin.strip())
    return "\n\n".join(sayfalar)


def metin_cikar(dosya_adi: str, veri: bytes) -> str:
    uzanti = Path(dosya_adi).suffix.lower()
    if uzanti == ".docx":
        return word_oku(veri)
    if uzanti == ".pdf":
        return pdf_oku(veri)
    if uzanti == ".txt":
        return veri.decode("utf-8", errors="replace")
    raise ValueError(f"Desteklenmeyen dosya turu: {uzanti} (.docx, .pdf veya .txt olmali)")


# ------------------------------------------------------- sayfa metinleri (Word)
_SAYFA_ISARETI = re.compile(
    r"^[\s\-=#*_>]*(?:sayfa|page)\s*[-#]?\s*(\d{1,4})\s*(?:[:.\-–—=#*_>)]+\s*(.*?))?[\s\-=#*_<]*$",
    re.IGNORECASE,
)


def _govde_bloklari(belge):
    """Word govdesindeki paragraf ve tablolari belgedeki sirasiyla dondurur.
    Doner: [(metin, sayfa_sonu_once, sayfa_sonu_sonra)]"""
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for el in belge.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, belge)
            once = p.paragraph_format.page_break_before or False
            sonra = any(br.get(qn("w:type")) == "page" for br in el.iter(qn("w:br")))
            yield p.text, once, sonra
        elif el.tag == qn("w:tbl"):
            for satir in Table(el, belge).rows:
                hucreler = []
                for h in satir.cells:
                    t = h.text.strip()
                    if t and t not in hucreler:   # birlesik hucreler tekrar etmesin
                        hucreler.append(t)
                if hucreler:
                    yield " | ".join(hucreler), False, False


def word_sayfa_metinleri(veri: bytes) -> tuple[dict[int, str], str]:
    """Word dosyasini PDF sayfalarina gore boler.

    1) "Sayfa 1", "Sayfa 2" ... satirlari varsa onlara gore (sayfa atlanabilir, sira serbest)
       Isaret satirinda metin de olabilir: "Sayfa 3: Kapanis metni..."
    2) Yoksa Word'deki sayfa sonlarina (Ctrl + Enter) gore sirayla
    Doner: ({sayfa_no (1'den): metin}, kullanilan_yontem)
    """
    from docx import Document

    bloklar = list(_govde_bloklari(Document(BytesIO(veri))))

    if any(_SAYFA_ISARETI.match(m.strip()) for m, _, _ in bloklar if m.strip()):
        sayfalar: dict[int, list[str]] = {}
        aktif = None
        for metin, _, _ in bloklar:
            metin = metin.strip()
            if not metin:
                continue
            m = _SAYFA_ISARETI.match(metin)
            if m:
                aktif = int(m.group(1))
                sayfalar.setdefault(aktif, [])
                if m.group(2):
                    sayfalar[aktif].append(m.group(2).strip())
            elif aktif is not None:   # ilk isaretten onceki metin (baslik vb.) yok sayilir
                sayfalar[aktif].append(metin)
        return {n: "\n".join(s).strip() for n, s in sayfalar.items()}, "isaret"

    sayfalar, aktif = {1: []}, 1
    for metin, once, sonra in bloklar:
        if once and any(sayfalar[aktif]):
            aktif += 1
            sayfalar[aktif] = []
        if metin.strip():
            sayfalar[aktif].append(metin.strip())
        if sonra:
            aktif += 1
            sayfalar[aktif] = []
    sonuc = {n: "\n".join(s).strip() for n, s in sayfalar.items()}
    while sonuc and not sonuc[max(sonuc)]:   # sondaki bos sayfa
        sonuc.pop(max(sonuc))
    return sonuc, "sayfa_sonu"


def sayfalara_esle(word_sayfalari: dict[int, str], pdf_sayfa_sayisi: int) -> tuple[list[str], list[str]]:
    """Word metinlerini PDF sayfa listesine yerlestirir. Doner: (metinler, uyarilar)"""
    metinler = [word_sayfalari.get(i + 1, "") for i in range(pdf_sayfa_sayisi)]
    uyarilar = []
    fazla = sorted(n for n in word_sayfalari if n > pdf_sayfa_sayisi or n < 1)
    if fazla:
        uyarilar.append(f"Word'de PDF'te olmayan sayfa numaraları var, yok sayıldı: {', '.join(map(str, fazla))} "
                        f"(PDF {pdf_sayfa_sayisi} sayfa)")
    bos = [str(i + 1) for i, m in enumerate(metinler) if not m]
    if bos:
        uyarilar.append(f"Metni olmayan sayfalar sessiz gösterilecek: {', '.join(bos)}")
    return metinler, uyarilar
